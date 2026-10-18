"""Describe the frozen JPCR patients and slides without changing predictions."""
from pathlib import Path
import hashlib
import json
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'rev'
OUT.mkdir(exist_ok=True)
CANCERS = dict(BRCA='Brca', HNSC='Hnscc', KIRC='Ccrcc', LUAD='Luad', LUSC='Lscc', OV='Ov', UCEC='Ucec')
tcga_path = ROOT / 'external_hrd_sources/pancancer_clinical_extension_20260612/TCGA-CDR-SupplementalTableS1.xlsx'
cptac_path = ROOT / 'external_hrd_sources/cptac_package/mssm-all_cancers-clinical-clinical_Pan-cancer.May2022.tsv.gz'
tcga = pd.read_excel(tcga_path, sheet_name='TCGA-CDR').set_index('bcr_patient_barcode', verify_integrity=True)
cptac = pd.read_csv(cptac_path, sep='\t').set_index('case_id', verify_integrity=True)
prov = ROOT / 'provenance/cptac_tumor_slide_metadata_20260719'
tcia_path = prov / 'cptac_metadata_07-09-2024.csv'
idc_path = prov / 'idc_v24_cptac_brca_ov_full_slide_mapping_20260726.csv'
tcia = pd.read_csv(tcia_path).set_index('Slide_ID')
idc = pd.read_csv(idc_path).set_index('FILENAME', verify_integrity=True)

def clean(v):
    if pd.isna(v) or str(v).startswith('[') or str(v).lower() in {'unknown','not reported','not applicable','not specified'}:
        return None
    return str(v)

def stage(v):
    v = clean(v)
    if not v: return 'Unknown'
    m = re.match(r'^(?:Stage\s*)?(IV|III|II|I|0)', v, re.I)
    return m.group(1).upper() if m else 'Unknown'

patients, slides, flow, eligibility = [], [], [], []
for cohort in ['TCGA', 'CPTAC']:
    for cancer, local in CANCERS.items():
        suffix = 'tcga_oof' if cohort == 'TCGA' else 'cptac_label_free'
        pred = pd.read_csv(ROOT / f'full210_equal_model_average_20260904/frozen_predictions/{cancer.lower()}_{suffix}.tsv', sep='\t')
        ids = set(pred.PATIENT)
        assert len(ids) == len(pred)
        base = f'tcga_sigprof_paper_sweep/by_config/legacy30_min50/{local}/TCGA_{cancer}' if cohort == 'TCGA' else f'cptac_sigprof_legacy30_min50/by_cancer/{local}/CPTAC_{local}'
        counts = pd.read_csv(ROOT / (base + '_legacy30_min50_sample_snv_count.csv'))
        raw = pd.read_excel(ROOT / (base + '_legacy30_min50_raw.xlsx'))
        eligible = set(raw.PATIENT)
        assert eligible == set(counts.loc[counts.snv_count > 50, 'PATIENT'])
        targets = pd.read_excel(ROOT / f'process_endpoint_inputs_legacy30_min50/{cohort}/{local}/{cohort}_{local}_process_endpoints.xlsx')
        prepared = set(targets.PATIENT)
        assert ids <= prepared <= eligible
        inventory = pd.read_csv(ROOT / f'sigprofiler_improved_20260818/inputs/nature_sigprofiler24_v1/matched/{cohort}/{cancer}/slides.tsv', sep='\t')
        assert set(inventory.PATIENT) == ids
        alignment = pd.read_csv(ROOT / f'legacy30_retccl_ctranspath_feature_fusion_20260829/fixed_bag_cache/{cohort}_{cancer}_alignment_audit.tsv', sep='\t').set_index('PATIENT')
        assert inventory.groupby('PATIENT').size().sort_index().equals(alignment.slides.sort_index())
        all_ids = set(counts.PATIENT)
        flow.append(dict(cohort=cohort,cancer=cancer,SNV_catalogue_samples=len(counts),SNV_catalogue_patients=len(all_ids),
                         excluded_no_sample_over_50=len(all_ids-eligible),eligible_genomic_samples=int((counts.snv_count>50).sum()),
                         eligible_genomic_patients=len(eligible),prepared_target_patients=len(prepared),
                         excluded_before_prepared_target=len(eligible-prepared),excluded_feature_match=len(prepared-ids),
                         total_excluded_after_genomic_eligibility=len(eligible-ids),analysed_patients=len(ids),analysed_slides=len(inventory)))
        for pid in sorted(all_ids):
            status = ('Included' if pid in ids else 'No sample with >50 observed substitutions' if pid not in eligible
                      else 'Absent from prepared model inputs' if pid not in prepared else 'No matched usable feature bag')
            eligibility.append(dict(cohort=cohort,cancer=cancer,PATIENT=pid,maximum_sample_SNV_count=int(counts.loc[counts.PATIENT==pid,'snv_count'].max()),status=status))
        for pid in sorted(ids):
            src = (tcga if cohort=='TCGA' else cptac).loc[pid]
            if cohort == 'TCGA':
                sf = 'clinical_stage' if cancer in ['OV','UCEC'] else 'ajcc_pathologic_tumor_stage'
                age,sex,raw_stage = src.age_at_initial_pathologic_diagnosis,src.gender,src[sf]
                year = pd.to_numeric(src.initial_pathologic_dx_year,errors='coerce')
                hist,grade,race = src.histological_type,src.histological_grade,src.race
            else:
                sf = 'baseline/tumor_stage_pathological'
                age,sex,raw_stage = src['consent/age'],src['consent/sex'],src[sf]
                year = None
                hist,grade,race = src['baseline/histologic_type'],src['cptac_path/histologic_grade'],src['consent/race']
            patients.append(dict(cohort=cohort,cancer=cancer,PATIENT=pid,age_years=pd.to_numeric(age,errors='coerce'),
                                 age_definition='At diagnosis' if cohort=='TCGA' else 'At consent',sex=clean(sex).title() if clean(sex) else None,
                                 stage_group=stage(raw_stage),stage_raw=None if pd.isna(raw_stage) else str(raw_stage),stage_source_field=sf,
                                 diagnosis_year=year,collection_date=None,histology=clean(hist),grade=clean(grade),race=clean(race),
                                 slide_count=int(alignment.loc[pid,'slides'])))
        for row in inventory.itertuples():
            x=dict(cohort=cohort,cancer=cancer,PATIENT=row.PATIENT,FILENAME=row.FILENAME,slide_category=None,embedding=None,
                   specimen_type=None,tumor_segment_acceptable=None,metadata_source=None)
            if cohort == 'TCGA':
                code=re.search(r'-((?:DX|TS|BS|MS)[A-Z0-9]+)(?:\.|$)',row.FILENAME).group(1)
                x.update(slide_category=code,embedding='FFPE diagnostic' if code.startswith('DX') else 'Frozen tissue section',
                         specimen_type='Primary tumour (sample type 01)',metadata_source='GDC barcode; Cooper et al. 2018')
            elif cancer in ['BRCA','OV']:
                m=idc.loc[row.FILENAME]
                typ={'HE tumor_tissue':'tumor_tissue','HE normal_tissue':'normal_tissue'}.get(m.dicom_SeriesDescription,'unspecified')
                x.update(slide_category=m.dicom_SeriesDescription,embedding='Not supplied in matched IDC extract',specimen_type=typ,metadata_source='IDC v24 DICOM metadata')
            else:
                m=tcia.loc[row.FILENAME]
                assert isinstance(m, pd.Series), 'Ambiguous metadata for an analysed slide'
                assert m.Case_ID == row.PATIENT
                x.update(slide_category='H&E',embedding=m.Embedding_Medium,specimen_type=m.Specimen_Type,
                         tumor_segment_acceptable=None if pd.isna(m.Tumor_Segment_Acceptable) else m.Tumor_Segment_Acceptable,
                         metadata_source='TCIA pathology metadata 2024-07-09')
            slides.append(x)

patients,slides,flow,eligibility = map(pd.DataFrame,[patients,slides,flow,eligibility])
assert len(patients)==3321 and len(slides)==7374
summary=[]
for (cohort,cancer),d in patients.groupby(['cohort','cancer'],sort=False):
    age=d.age_years.dropna(); years=d.diagnosis_year.dropna()
    s=dict(cohort=cohort,cancer=cancer,patients=len(d),age_median=age.median(),age_Q1=age.quantile(.25),age_Q3=age.quantile(.75),
           age_min=age.min(),age_max=age.max(),age_missing=int(d.age_years.isna().sum()),female=int((d.sex=='Female').sum()),
           male=int((d.sex=='Male').sum()),sex_missing=int(d.sex.isna().sum()),stage_I=int((d.stage_group=='I').sum()),
           stage_II=int((d.stage_group=='II').sum()),stage_III=int((d.stage_group=='III').sum()),stage_IV=int((d.stage_group=='IV').sum()),
           stage_other_unknown=int((~d.stage_group.isin(['I','II','III','IV'])).sum()),
           diagnosis_year_min=years.min(),diagnosis_year_max=years.max(),diagnosis_year_missing=int(d.diagnosis_year.isna().sum()),
           age_definition=d.age_definition.iloc[0],stage_source_field=d.stage_source_field.iloc[0])
    summary.append(s)
summary=pd.DataFrame(summary)
categories=patients.melt(id_vars=['cohort','cancer','PATIENT'],value_vars=['histology','grade','race','stage_raw'],var_name='characteristic',value_name='category').fillna({'category':'Not reported'})
categories=categories.groupby(['cohort','cancer','characteristic','category'],dropna=False).size().reset_index(name='patients')
slide_summary=slides.groupby(['cohort','cancer','embedding','specimen_type'],dropna=False).agg(slides=('FILENAME','size'),patients=('PATIENT','nunique')).reset_index()

tables={'cohort_summary':summary,'patient_characteristics':patients,'patient_flow':flow,'eligibility_audit':eligibility,
        'slide_summary':slide_summary,'slide_manifest':slides,'clinical_categories':categories}
for name,df in tables.items():df.to_csv(OUT/f'{name}.tsv',sep='\t',index=False)
stats={'flow_totals':flow.groupby('cohort').sum(numeric_only=True).to_dict('index'),
       'embedding_counts':slides.groupby(['cohort','embedding']).size().to_dict(),
       'specimen_counts':slides.groupby(['cohort','specimen_type']).size().to_dict(),
       'source_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [tcga_path,cptac_path,tcia_path,idc_path]}}
stats['embedding_counts']={'|'.join(k):int(v) for k,v in stats['embedding_counts'].items()}
stats['specimen_counts']={'|'.join(k):int(v) for k,v in stats['specimen_counts'].items()}
(OUT/'cohort_audit.json').write_text(json.dumps(stats,indent=2))
print(summary.to_string(index=False))
print(json.dumps(stats,indent=2))
