"""Complete the supplementary captions without changing the plotted results."""
from pathlib import Path
import sys
from s1_rows import TITLE, CAPTION

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'Revision_Cancers'
LEGENDS = [
    (TITLE + '.', CAPTION),
    (
        'Figure S2. Target distributions, sparsity and catalogue counts.',
        'Observed SBS1–SBS30 exposures and standardised model scores are shown for TCGA (left two columns) and CPTAC (right two columns), with one cancer type per row. Boxes span the interquartile range, centre lines mark medians and whiskers extend to the 5th and 95th percentiles. Outliers are omitted. Yellow-to-brown strips show the fraction of patients with zero exposure for each signature. n denotes patients. Samples denotes included eligible genomic samples. Single-nucleotide variant (SNV) counts are sample-level medians [interquartile ranges]. Complete distributions, non-zero counts, mutation counts and eligibility summaries are provided in Supplementary Data S3 (distributions and mutation_summary).',
    ),
    (
        'Figure S3. Robustness of score–exposure associations.',
        '(a–b) Pearson r versus Spearman ρ in TCGA and CPTAC. Each point represents one cancer–signature pair, coloured by cancer type. Dashed diagonals mark equality. Effect agreement is the Pearson correlation between these two sets of coefficients across pairs. (c) The original 14 nominally replicated pairs under all-slide inputs (grey) and tumour-labelled-slide inputs (teal), evaluated in the same 570 CPTAC patients. Points show Pearson r and bars show Fisher-z 95% confidence intervals. The dashed line marks zero. Restriction uses source-slide annotations, not tile-level tumour classification. Ten patients lacking tumour-labelled inputs were excluded only from this comparison. The primary analysis retains 580 patients. Complete 210-pair sensitivity results and input counts are provided in Supplementary Data S3.',
    ),
    (
        'Figure S4. Exploratory genomic-feature associations and MANTIS sensitivity.',
        '(a) External discrimination in 82 CPTAC UCEC patients using the primary SBS6 score for microsatellite instability (MSI, 25 positive cases) and the SBS10 score for the predefined POLE/POLD1 composite (six positive cases). P values are unadjusted, two-sided Mann–Whitney U tests. (b) TCGA MSI discrimination using the same SBS6 scores under the primary MANTIS rule (<0.4 negative, >0.6 positive, intermediate values excluded) and a 0.4 threshold (≤0.4 negative, >0.4 positive). In a–b, points show areas under the receiver operating characteristic curve (AUCs), bars show 95% confidence intervals from 10,000 class-stratified patient-bootstrap replicates and dashed lines mark AUC=0.5. (c) TCGA score distributions under the primary rule, with negatives in grey and positives in teal. (d) Receiver operating characteristic curves for the primary TCGA rule (teal), the TCGA 0.4 threshold (ochre) and CPTAC MSI labels (grey). The diagonal marks chance discrimination. The same primary SBS6 score is used in the signature analysis and MSI assessments. MSI labels were used only for evaluation. Endpoint definitions and results are provided in Supplementary Data S2 and S3.',
    ),
    (
        'Figure S5. Paired TCGA/CPTAC attention examples.',
        'Eight illustrative cases cover UCEC–SBS13, UCEC–SBS3, BRCA–SBS3 and LUAD–SBS4, with TCGA on the left and CPTAC on the right. Cases were selected near the median positive genomic exposure among patients with tumour-labelled slides. Each case shows its exposure, standardised ensemble score, H&E overview, attention overlay and three higher- and three lower-weight tiles. Colours indicate within-patient attention percentile in one MIL component, from dark for lower weights to yellow for higher weights. Uncoloured regions were not sampled. Tile labels show component attention weights. Each crop spans 256 × 256 µm, with an 80 µm scale bar. These examples locate regions for inspection but do not establish tissue identity, morphological mechanisms or contributions to the complete ensemble. Expert pathological review was not performed.',
    ),
]

S5_CAPTIONS = [LEGENDS[4][1]] + [
    f'{target} examples from TCGA ({tcga}, left) and CPTAC ({cptac}, right). '
    'Display conventions and case-selection criteria are described on the first page of Figure S5.'
    for target, tcga, cptac in [
        ('UCEC–SBS3', 'TCGA-A5-A0GV', 'C3N-01346'),
        ('BRCA–SBS3', 'TCGA-BH-A0C0', '01BR042'),
        ('LUAD–SBS4', 'TCGA-L9-A443', 'C3N-00704'),
    ]
]


def update_pdf():
    import fitz
    path = OUT / 'Supplementary_Figures.pdf'
    doc = fitz.open(path)
    assert len(doc) == 8
    font_path = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    font = fitz.Font(fontfile=font_path)
    for i, page in enumerate(doc):
        if i == 0:
            # s1_rows.py places the current legend within its A4 layout.
            continue
        caption = LEGENDS[i][1] if i < 4 else S5_CAPTIONS[i - 4]
        blocks = [b for b in page.get_text('blocks') if b[6] == 0 and b[4].strip()]
        old = max(blocks, key=lambda b: b[3])
        assert old[1] > page.rect.height * .90
        plot_bottom = max(b[3] for b in blocks if b is not old)
        left, right = page.rect.width * .045, page.rect.width * .975
        size, line_height = 9, 11.4
        lines, line = [], ''
        for word in caption.split():
            candidate = (line + ' ' + word).strip()
            if font.text_length(candidate, fontsize=size) > right - left:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
        last_baseline = page.rect.height - 12
        first_baseline = last_baseline - (len(lines) - 1) * line_height
        caption_top = first_baseline - font.ascender * size
        assert caption_top >= plot_bottom + 7, (i + 1, caption_top, plot_bottom)
        # Remove only caption text, preserving plot vectors and source images.
        page.add_redact_annot(fitz.Rect(old[:4]) + (-.5, -.5, .5, .5), fill=False)
        page.apply_redactions(images=0, graphics=0)
        page.insert_font(fontname='LegendFont', fontfile=font_path)
        for j, line in enumerate(lines):
            page.insert_text((left, first_baseline + j * line_height), line,
                             fontsize=size, fontname='LegendFont', color=(0, 0, 0))
        print(f'Page {i + 1}: {len(lines)} legend lines; gap {caption_top - plot_bottom:.1f} pt')
    temp = path.with_suffix('.tmp.pdf')
    doc.save(temp, garbage=4, deflate=True)
    doc.close()
    temp.replace(path)


def write_word():
    from docx import Document
    from docx.shared import Inches, Pt
    import subprocess
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.05
    for title, caption in LEGENDS:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.add_run(title).bold = True
        p = doc.add_paragraph(caption)
        p.paragraph_format.keep_together = True
    doc.save(OUT / 'Supplementary_Legends.docx')
    subprocess.run(['/opt/anaconda/envs/cenv/bin/python', __file__, '--pdf'], check=True)


if __name__ == '__main__':
    update_pdf() if '--pdf' in sys.argv else write_word()
