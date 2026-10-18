"""Prepare readable supplementary tables from the existing, verified values."""
from pathlib import Path
from copy import copy
import json
import math
import sys
import textwrap

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'Revision_Cancers'
BLUE, PALE = '24475A', 'F2F6F8'
P_FORMAT = '[<0.001]0.00E+00;0.000'
CANCERS = ['BRCA', 'HNSC', 'KIRC', 'LUAD', 'LUSC', 'OV', 'UCEC']
LABELS = {
    'cohort': 'Cohort', 'cancer': 'Cancer type', 'signature': 'SBS signature',
    'n': 'Patients (n)', 'nonzero': 'Non-zero values (n)', 'zero_fraction': 'Zero values (%)',
    'r': 'Pearson r', 'p': 'P (one-sided)', 'ci_low': '95% CI lower', 'ci_high': '95% CI upper',
    'direction': 'Correlation directions', 'nominal_positive_both': 'Nominally replicated\nassociation',
    'all_n': 'All-slide inputs\nPatients (n)', 'all_r': 'All-slide inputs\nPearson r',
    'all_p': 'All-slide inputs\nP (one-sided)', 'all_ci_low': 'All-slide inputs\n95% CI lower',
    'all_ci_high': 'All-slide inputs\n95% CI upper', 'tumour_n': 'Tumour-labelled inputs\nPatients (n)',
    'tumour_r': 'Tumour-labelled inputs\nPearson r', 'tumour_p': 'Tumour-labelled inputs\nP (one-sided)',
    'tumour_q_bh210': 'Tumour-labelled inputs\nBH-adjusted P',
    'tumour_ci_low': 'Tumour-labelled inputs\n95% CI lower',
    'tumour_ci_high': 'Tumour-labelled inputs\n95% CI upper',
    'delta_r': 'Change in Pearson r', 'score_correlation': 'Correlation between input-condition scores',
    'endpoint': 'Endpoint', 'rule': 'Classification rule', 'positive_n': 'Positive patients (n)',
    'auc': 'ROC AUC', 'variable': 'Measurement', 'minimum': 'Minimum', 'maximum': 'Maximum',
    'median': 'Median', 'q1': '25th percentile', 'q3': '75th percentile', 'mean': 'Mean', 'sd': 'Standard deviation',
    'source_samples': 'Source genomic samples (n)', 'source_patients': 'Source patients (n)',
    'eligible_genomic_samples': 'Genomic samples with >50 SNVs (n)',
    'eligible_genomic_patients': 'Genomically eligible patients (n)',
    'included_patients': 'Analysed patients (n)', 'included_eligible_samples': 'Analysed genomic samples (n)',
    'embedding': 'Specimen preparation', 'specimen_type': 'Source-slide tissue label',
    'slides': 'Slides (n)', 'patients': 'Patients (n)', 'stratum': 'Patient input group',
    'method': 'Adjustment method', 'error_rate': 'Error rate addressed',
    'adjusted_P_threshold': 'Adjusted P threshold (<)', 'TCGA_pairs': 'TCGA pairs below threshold (n)',
    'CPTAC_pairs': 'CPTAC pairs below threshold (n)',
    'pairs_below_threshold_in_both': 'Pairs below threshold in both cohorts (n)',
    'original14_below_threshold_in_both': 'Original 14 below threshold in both cohorts (n)',
    'original14': 'Nominally replicated\nassociation', 'matched_patients': 'Matched patients (n)',
    'status': 'Comparison availability', 'original_nonzero': 'Original exposure\nNon-zero patients (n)',
    'alternative_nonzero': 'Alternative exposure\nNon-zero patients (n)',
    'exposure_r': 'Original vs alternative exposures\nPearson r',
    'exposure_rho': 'Original vs alternative exposures\nSpearman ρ',
    'mean_absolute_difference': 'Original vs alternative exposures\nMean absolute difference',
    'original_r': 'Original exposure\nScore Pearson r',
    'original_p_one_sided': 'Original exposure\nP (one-sided)',
    'original_ci_low': 'Original exposure\n95% CI lower', 'original_ci_high': 'Original exposure\n95% CI upper',
    'alternative_r': 'Alternative exposure\nScore Pearson r',
    'alternative_p_one_sided': 'Alternative exposure\nP (one-sided)',
    'alternative_ci_low': 'Alternative exposure\n95% CI lower', 'alternative_ci_high': 'Alternative exposure\n95% CI upper',
    'original_q_bh210': 'Original exposure\nBH-adjusted P',
    'alternative_q_bh210': 'Alternative exposure\nBH-adjusted P',
    'analysis': 'Analysis', 'positives': 'Positive patients (n)', 'metric': 'Metric', 'estimate': 'Estimate',
    'p_value': 'P (two-sided)', 'label_source': 'Label source',
    'exclusion_or_intersection': 'Eligibility / exclusions', 'included_n': 'Included patients (n)',
    'PATIENT': 'Patient ID', 'named_signature': 'SBS signature', 'score': 'Image-derived score',
    'observed': 'Observed SBS exposure (proportion)', 'label': 'Endpoint label (0/1)',
    'hrd_score': 'Genomic-scar HRD score', 'mantis_score': 'MANTIS score',
    'pole_positive': 'Qualifying POLE variant (0/1)', 'pold1_positive': 'Qualifying POLD1 variant (0/1)',
    'gene': 'Gene', 'protein_change': 'Protein change', 'hgvsc': 'Coding DNA change (HGVS)',
    'chromosome': 'Chromosome', 'position_grch37': 'Genomic position (GRCh37)',
}
for co in ['TCGA', 'CPTAC']:
    for key, label in {
        'n': 'Patients (n)', 'nonzero': 'Non-zero exposure (n)', 'pearson_r': 'Pearson r',
        'ci_low': '95% CI lower', 'ci_high': '95% CI upper',
        'pearson_p_positive_direction': 'P (one-sided)', 'q_bh210': 'BH-adjusted P',
        'spearman_rho': 'Spearman ρ', 'spearman_p_two_sided': 'Spearman P (two-sided)',
    }.items():
        LABELS[key + '_' + co] = co + '\n' + label


def records(sheet, header):
    fields = [c.value for c in sheet[header]]
    return fields, [dict(zip(fields, row)) for row in sheet.iter_rows(min_row=header + 1, values_only=True)]


def build(book, name, title, note, fields, rows, overrides=None, freeze='C5'):
    labels = dict(LABELS)
    labels.update(overrides or {})
    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:4'
    sheet.print_options.horizontalCentered = True
    sheet.freeze_panes = freeze
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(fields))
    sheet['A1'] = title
    sheet['A1'].font = Font(name='Arial', size=14, bold=True, color=BLUE)
    sheet['A1'].alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 26
    for j, field in enumerate(fields, 1):
        label = labels[field]
        width = 16
        if field in ['cancer','cohort','signature','gene','chromosome','named_signature']: width = 12
        if field in ['direction','method','error_rate','PATIENT','endpoint','embedding','specimen_type','stratum']: width = 26
        if field in ['analysis','rule','label_source','exclusion_or_intersection','status']: width = 42
        if field == 'hgvsc': width = 32
        if name == 'polymerase_rules': width = 18 if field == 'gene' else 40
        sheet.column_dimensions[get_column_letter(j)].width = width
        cell = sheet.cell(4, j, label)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.comment = Comment('Analysis field: ' + field, 'Study authors')
    total_width = sum(sheet.column_dimensions[get_column_letter(j)].width for j in range(1,len(fields)+1))
    sheet.row_dimensions[1].height = max(26, math.ceil(len(title)/(total_width*.65))*18+8)
    note_end = min(len(fields), 8)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=note_end)
    sheet['A2'] = note
    sheet['A2'].font = Font(name='Arial', size=10, color='333333')
    sheet['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    note_width = int(sum(sheet.column_dimensions[get_column_letter(j)].width for j in range(1, note_end+1)) * 1.05)
    lines = sum(max(1, len(textwrap.wrap(p, max(note_width, 40)))) for p in note.split('\n'))
    sheet.row_dimensions[2].height = max(30, lines * 14 + 8)
    sheet.row_dimensions[3].height = 7
    sheet.row_dimensions[4].height = 58
    for i, row in enumerate(rows, 5):
        max_lines = 1
        for j, field in enumerate(fields, 1):
            v = row[field]
            cell = sheet.cell(i, j, v)
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='center', wrap_text=True,
                                       horizontal='right' if isinstance(v, (int,float)) and not isinstance(v,bool) else 'left')
            if i % 2 == 1: cell.fill = PatternFill('solid', fgColor=PALE)
            label = labels[field]
            if isinstance(v, bool): cell.alignment = Alignment(horizontal='center', vertical='center')
            elif isinstance(v, (int,float)):
                if '(n)' in label or '(0/1)' in label or field in ['chromosome','position_grch37']:
                    cell.number_format = '#,##0'
                elif field == 'zero_fraction': cell.number_format = '0.0%'
                elif field == 'adjusted_P_threshold': cell.number_format = '0.00'
                elif field in ['p','p_value'] or '_p' in field or field.startswith('q_') or '_q_' in field:
                    cell.number_format = P_FORMAT
                else: cell.number_format = '0.000'
            if isinstance(v, str):
                width = sheet.column_dimensions[get_column_letter(j)].width
                max_lines = max(max_lines, sum(max(1,len(textwrap.wrap(s, max(int(width),8)))) for s in v.split('\n')))
        sheet.row_dimensions[i].height = max(21, max_lines * 13 + 5)
    sheet.auto_filter.ref = f'A4:{get_column_letter(len(fields))}{sheet.max_row}'
    sheet.print_area = f'A1:{get_column_letter(len(fields))}{sheet.max_row}'
    # A reformatted table must preserve every retained underlying value.
    for i, row in enumerate(rows, 5):
        assert [sheet.cell(i,j).value for j in range(1,len(fields)+1)] == [row[f] for f in fields]
    return sheet


def sort_rows(rows):
    return sorted(rows, key=lambda r:(r.get('method',''), ['TCGA','CPTAC'].index(r['cohort']) if r.get('cohort') in ['TCGA','CPTAC'] else 0,
        CANCERS.index(r['cancer']) if r.get('cancer') in CANCERS else 0,
        int(r['signature'][3:]) if str(r.get('signature','')).startswith('SBS') else 0,
        r.get('variable',''), r.get('stratum','')))


def main():
    before = {n:load_workbook(OUT/f'Supplementary_Data_S{n}.xlsx') for n in [1,2,3]}
    s1,s2,s3 = (before[n] for n in [1,2,3])
    # The full S3 table contains every numeric result in the duplicated S1 table.
    _, pair_rows = records(s3['paired_metrics'],1)
    pair_map = {(r['cancer'],r['signature']):r for r in pair_rows}
    _, repeated = records(s1['210_pair_results'],4)
    mapping = {'TCGA patients (n)':'n_TCGA','TCGA non-zero exposure (n)':'nonzero_TCGA',
        'TCGA Pearson r':'pearson_r_TCGA','TCGA P (one-sided)':'pearson_p_positive_direction_TCGA',
        'TCGA Spearman ρ':'spearman_rho_TCGA','CPTAC patients (n)':'n_CPTAC',
        'CPTAC non-zero exposure (n)':'nonzero_CPTAC','CPTAC Pearson r':'pearson_r_CPTAC',
        'CPTAC P (one-sided)':'pearson_p_positive_direction_CPTAC','CPTAC Spearman ρ':'spearman_rho_CPTAC'}
    for row in repeated:
        target=pair_map[row['Cancer type'],row['SBS signature']]
        assert all(row[k] == target[v] for k,v in mapping.items())
    del s1['210_pair_results']
    s1.active=0
    for sheet in s1:
        sheet.sheet_view.showGridLines=False
        sheet.sheet_view.zoomScale=90
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape'
        sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1
        sheet.page_setup.fitToHeight=0
        sheet.print_area=sheet.dimensions
    # Correct citation numbers against the current manuscript reference list.
    details=s1['cohort_details']
    for row in details:
        if row[0].value=='TCGA clinical characteristics': row[6].value=row[6].value.replace('[37]','[38]')
        if row[0].value=='TCGA slide interpretation': row[6].value=row[6].value.replace('[38]','[39]')
        if row[0].value=='TCGA slide annotations': row[1].value=row[1].value.replace('[38]','[39]')
    s1['model_settings']['A2'] = s1['model_settings']['A2'].value + ' Complete 210-pair association results are in Data S3 (paired_metrics).'
    s1['model_settings'].row_dimensions[2].height=80
    s1.save(OUT/'Supplementary_Data_S1.xlsx')

    new=Workbook();new.remove(new.active)
    removed_events=0
    for sheet in s2:
        fields,rows=records(sheet,4)
        note=sheet['A2'].value
        if sheet.title=='patient_scores':
            note=note.replace('score is','Image-derived score is').replace('observed is','Observed SBS exposure is').replace('label is','Endpoint label is')
        if sheet.title=='polymerase_events':
            removed_events=sum(r['analysis_cohort_included']==0 for r in rows)
            rows=[r for r in rows if r['analysis_cohort_included']==1]
            assert len(rows)==49 and len({r['PATIENT'] for r in rows})==47
            fields.remove('analysis_cohort_included')
            note='The 49 qualifying variant events observed in 47 mutation-positive patients within the 422-patient TCGA UCEC analysis. Patients with more than one event occur on multiple rows. Coordinates use GRCh37. HGVS denotes Human Genome Variation Society nomenclature. The qualifying substitution list is in polymerase_rules.'
        labels={'p':'P (two-sided)'} if sheet.title=='clinical_results' else {}
        build(new,sheet.title,sheet['A1'].value,note,fields,rows,labels,freeze='D5' if sheet.title=='patient_scores' else 'B5')
    new.save(OUT/'Supplementary_Data_S2.xlsx')

    notes={
        'paired_metrics':('Complete 210-pair association results',
            'All seven cancer types × SBS1–SBS30. TCGA and CPTAC results are grouped separately. CI: 95% Fisher-z confidence interval. Pearson P values test a positive association (one-sided); Spearman P values are two-sided. BH: Benjamini–Hochberg adjustment across 210 tests separately in each cohort. “Original 14 pairs?” identifies positive Pearson associations with unadjusted P<0.05 in both cohorts. Blank correlations reflect constant exposure or score values; these tests entered adjustment as P=1. Full-precision values are retained.'),
        'tumour_sensitivity':('Paired tumour-labelled-slide input sensitivity',
            'Across seven cancer types, 570 CPTAC patients were compared using both input conditions: 2,221 all-input slides versus 1,509 tumour-labelled slides. Ten patients lacking tumour-labelled inputs were excluded only here; the primary cohort remains 580. Restriction uses source-slide annotations. CI: 95% Fisher-z interval. P values are one-sided; BH adjustment covers all 210 restricted-input tests. Change in r = tumour-labelled minus all-slide r. The final column correlates the two sets of image scores. Blank statistics indicate constant values.'),
        'clinical_revision':('Exploratory external and MANTIS-sensitivity results',
            'Figure S4. MSI uses the final continuous-SBS6-trained clinical score; polymerase status uses the primary SBS10 score and the same qualifying POLE OR POLD1 list as Data S2. The primary MANTIS rule is <0.4 negative/>0.6 positive (0.4–0.6 excluded); the alternative is ≤0.4 negative/>0.4 positive. CPTAC labels are source-published MSI-H/MSS. AUC: area under the ROC curve; CIs use 10,000 class-stratified patient-bootstrap replicates. P values are unadjusted two-sided Mann–Whitney U tests. All analyses are exploratory.'),
        'distributions':('SBS exposure and model-score distributions',
            'Figure S2. Each row summarises one cohort, cancer type, SBS signature and measurement. Exposure is the observed genomic proportion (0–1); Score is the standardised image-derived score. Non-zero counts and zero percentages apply to the indicated measurement. Percentages are displayed from stored fractions (0–1). Quartiles and medians are patient-level summaries; standard deviations use the population definition (ddof=0).'),
        'mutation_summary':('Genomic sample eligibility and mutation counts',
            'Figure S2. Samples are genomic samples; patients are distinct individuals. Genomic eligibility requires >50 observed substitutions in at least one sample. SNV summaries describe included eligible genomic samples, not patient-averaged counts. Full patient-selection steps and exclusions are in Data S1 (cohort_details). SNV: single-nucleotide variant. Quartiles may be fractional because of interpolation.'),
        'slide_summary':('Image-input composition by source annotation',
            'Primary input slides grouped by cohort, cancer, preparation and source-slide tissue label. Slide categories are disjoint; patients can appear in several rows and must not be summed across categories. Tumour/normal labels describe source metadata, not every tile. FFPE: formalin-fixed paraffin-embedded; OCT: optimal cutting temperature embedding for frozen sections. Missing preparation annotations remain unspecified. Overall patient counts are in Data S1 (cohort_summary).'),
        'specimen_strata':('Descriptive associations by input composition',
            'Existing all-slide scores are stratified by patient input composition; models were not rerun for each specimen group. “Mixed” means both FFPE and frozen/OCT inputs; “Unspecified” means preparation annotations were unavailable. Normal-slide inclusion groups form a separate comparison and overlap the preparation groups. P values test positive Pearson correlation (one-sided, unadjusted); CIs are 95% Fisher-z intervals. Blank statistics indicate insufficient patients or no variation. Small strata are descriptive.'),
        'signature_comparison':('Sensitivity to signature-fitting method and reference',
            'Image scores are fixed; original and alternative exposures are compared in the same patients within each method/cohort/cancer. Matched patients counts shared IDs; Patients (n) counts available paired targets. Exposure comparisons use proportions (0–1); change in r = alternative minus original score–exposure r. CIs are 95% Fisher-z intervals. P values are one-sided; BH adjustment covers 210 tests per method/cohort/target definition. Blank statistics follow the stated availability; no subtype aggregation was used. These comparisons do not rank target accuracy.\n'),
    }
    method_notes={r['method']:r['source_definition'] for r in records(s3['signature_comparison'],1)[1]}
    notes['signature_comparison']=(notes['signature_comparison'][0],notes['signature_comparison'][1]+'\n'.join(f'{k}: {v.replace("corrected TCGA", "TCGA")}' for k,v in method_notes.items()))
    new=Workbook();new.remove(new.active)
    for sheet in s3:
        name=sheet.title
        fields,rows=records(sheet,1)
        if name=='paired_metrics':
            fields=['cancer','signature']+[key+'_'+co for co in ['TCGA','CPTAC'] for key in ['n','nonzero','pearson_r','ci_low','ci_high','pearson_p_positive_direction','q_bh210','spearman_rho','spearman_p_two_sided']]+['direction','nominal_positive_both']
        elif name=='tumour_sensitivity':
            fields=['cancer','signature','all_n','all_r','all_ci_low','all_ci_high','all_p','tumour_n','tumour_r','tumour_ci_low','tumour_ci_high','tumour_p','tumour_q_bh210','delta_r','score_correlation']
        elif name=='mutation_summary':
            assert all(r['samples_over50']==r['eligible_genomic_samples'] for r in rows)
            fields=['cohort','cancer','source_patients','source_samples','eligible_genomic_patients','eligible_genomic_samples','included_patients','included_eligible_samples','minimum','q1','median','q3','maximum']
        elif name=='specimen_strata': fields=['cohort','cancer','signature','stratum','n','r','ci_low','ci_high','p']
        elif name=='signature_comparison':
            fields=['method','cohort','cancer','signature','original14','matched_patients','n','status','original_nonzero','alternative_nonzero','exposure_r','exposure_rho','mean_absolute_difference','original_r','original_ci_low','original_ci_high','original_p_one_sided','original_q_bh210','alternative_r','alternative_ci_low','alternative_ci_high','alternative_p_one_sided','alternative_q_bh210','delta_r']
        title,note=notes[name]
        labels={}
        if name=='clinical_revision':labels={'p':'P (two-sided, unadjusted)'}
        if name=='signature_comparison': labels={'method':'Fitting method / reference','n':'Paired-target patients (n)'}
        if name=='mutation_summary':labels.update({f:'SNVs per sample\n'+LABELS[f] for f in ['minimum','q1','median','q3','maximum']})
        rows=rows if name=='clinical_revision' else sort_rows(rows)
        freeze='E5' if name=='signature_comparison' else 'D5' if name in ['distributions','specimen_strata'] else 'C5'
        result=build(new,name,'Supplementary Data S3. '+title,note,fields,rows,labels,freeze)
        if name=='mutation_summary':
            for j,field in enumerate(fields,1):
                if field in ['minimum','maximum','q1','median','q3']:
                    for row in result.iter_rows(min_row=5,min_col=j,max_col=j):
                        row[0].number_format='#,##0' if field in ['minimum','maximum'] else '#,##0.00'
    new.save(OUT/'Supplementary_Data_S3.xlsx')
    report={
        'S1_sheets':s1.sheetnames,'S2_sheets':[s.title for s in load_workbook(OUT/'Supplementary_Data_S2.xlsx')],
        'S3_sheets':new.sheetnames,'duplicate_210_pair_sheet_removed_after_value_match':True,
        'out_of_analysis_polymerase_events_removed':removed_events,
        'all_retained_values_preserved_before_save':True,
        'source_fields_in_header_comments':True,
        'visible_sheet_titles_definitions_and_statistics':True,
        'numeric_precision_only_changes_display':True,
        'source_citation_numbers_updated_to_manuscript':{'TCGA_CDR':38,'TCGA_pathology':39},
        'S3_repeated_metadata_moved_to_visible_notes':True,
    }
    (ROOT/'rev/xlsx_check.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))
    polish_sheets()


def polish_sheets():
    """Shorten reader-facing notes and labels without changing numerical results."""
    book = load_workbook(OUT / 'Supplementary_Data_S3.xlsx')
    before = {s.title: {(c.row, c.column): c.value for row in s for c in row}
              for s in book}
    notes = {
        'paired_metrics': (
            'CI: Fisher-z 95% confidence interval. BH: Benjamini–Hochberg adjustment across 210 tests within each cohort.\n'
            'Nominally replicated associations have positive Pearson correlations with one-sided P < 0.05 in both cohorts.\n'
            'Blank correlations indicate constant exposure or score values. These tests were included in BH adjustment with P = 1.'),
        'tumour_sensitivity': (
            'Paired comparison of all-slide and tumour-labelled-slide inputs in 570 CPTAC patients. The primary analysis retains 580 patients.\n'
            'Inputs comprise 2,221 and 1,509 source slides, respectively. Restriction uses slide annotations. Change in r = tumour-labelled minus all-slide Pearson r.\n'
            'CI: Fisher-z 95% confidence interval. P values are one-sided. BH adjustment covers 210 restricted-input tests. Blank statistics indicate constant values.'),
        'clinical_revision': (
            'Exploratory MSI and POLE/POLD1 analyses (Figure S4). MSI uses the exposure-trained SBS6 score. POLE/POLD1 uses the primary SBS10 score and the variant list in Data S2.\n'
            'Primary MANTIS rule: <0.4 negative, >0.6 positive; 0.4–0.6 excluded. Alternative rule: ≤0.4 negative, >0.4 positive. CPTAC uses published MSI-H/MSS labels.\n'
            'AUC: area under the ROC curve. CIs use 10,000 class-stratified bootstrap replicates. P values are unadjusted, two-sided Mann–Whitney U tests.'),
        'distributions': (
            'Patient-level distributions of genomic SBS exposure (proportion) and standardised image-derived scores (Figure S2).\n'
            'Non-zero counts and zero percentages refer to the measurement in each row. Standard deviations use the population definition.'),
        'mutation_summary': (
            'SNV: single-nucleotide variant. Genomic eligibility requires >50 substitutions in at least one sample.\n'
            'Mutation counts describe included genomic samples; patients and samples are counted separately. Cohort selection is detailed in Data S1.'),
        'slide_summary': (
            'Slides are grouped by cohort, cancer type, preparation and source-slide tissue label. Patient counts overlap across categories.\n'
            'Tissue labels refer to slide metadata. FFPE: formalin-fixed paraffin-embedded; OCT: optimal cutting temperature compound. Unspecified denotes missing annotation.'),
        'specimen_strata': (
            'Existing all-slide scores grouped by patient input composition. Predictions were not recalculated for each group. Mixed denotes FFPE and frozen/OCT inputs.\n'
            'Normal-slide inclusion and specimen preparation are separate, overlapping classifications. Small groups are descriptive.\n'
            'CI: Fisher-z 95% confidence interval. P values are one-sided and unadjusted. Blank statistics indicate insufficient patients or constant values.'),
        'signature_comparison': (
            'Fixed image-derived scores were compared with original and alternative exposures in the same patients within each method and cohort. Change in r = alternative minus original Pearson r.\n'
            'CI: Fisher-z 95% confidence interval. P values are one-sided. BH adjustment covers 210 tests per method, cohort and exposure definition.\n'
            'Only matching signature names were compared. Split signatures were not combined. Nominal replication denotes positive associations with P < 0.05 in both cohorts in the primary analysis.'),
    }
    for name, note in notes.items():
        sheet = book[name]
        sheet['A2'] = note
        end = min(sheet.max_column, 8)
        width = int(sum(sheet.column_dimensions[get_column_letter(j)].width
                        for j in range(1, end+1)) * 1.05)
        lines = sum(max(1, len(textwrap.wrap(p, width))) for p in note.split('\n'))
        sheet.row_dimensions[2].height = lines * 14 + 8
    flag_columns = {s.title: c.column for s in book for c in s[4]
                    if c.value in ('Nominally replicated\nassociation', 'Original 14 pairs?')}
    for name, column in flag_columns.items():
        sheet = book[name]
        sheet.cell(4, column, 'Nominally replicated\nassociation')
        sheet.column_dimensions[get_column_letter(column)].width = 24
        sheet.cell(4, column).comment = Comment(
            'Yes: positive Pearson correlations with unadjusted one-sided P < 0.05 in both TCGA and CPTAC in the primary analysis.',
            'Study authors')
        for row in range(5, sheet.max_row+1):
            cell = sheet.cell(row, column)
            if isinstance(cell.value, bool):
                cell.value = 'Yes' if cell.value else 'No'
            assert cell.value in ('Yes', 'No')
    methods = {
        'sigfit_legacy30': 'sigfit / legacy SBS1–SBS30',
        'SPA_v2': 'SigProfilerAssignment / COSMIC v2',
        'SPA_v3.5': 'SigProfilerAssignment / COSMIC v3.5',
        'SPA_v3.6': 'SigProfilerAssignment / COSMIC v3.6',
    }
    statuses = {
        'estimable': 'Available',
        'constant or insufficient values': 'Constant values or insufficient data',
        'signature not in source; no subtype aggregation': 'Signature unavailable',
    }
    sheet = book['signature_comparison']
    sheet.column_dimensions['A'].width = 35
    status_column = next(c.column for c in sheet[4] if c.value == 'Comparison availability')
    sheet.cell(4, status_column).comment = Comment(
        'Signature unavailable: no matching-name signature was supplied by the source. Split signatures were not combined.',
        'Study authors')
    for row in range(5, sheet.max_row+1):
        for column, mapping in [(1, methods), (status_column, statuses)]:
            cell = sheet.cell(row, column)
            cell.value = mapping.get(cell.value, cell.value)
        sheet.row_dimensions[row].height = max(34, sheet.row_dimensions[row].height or 0)
    # Every data-cell change is a defined display-label replacement; all numbers stay exact.
    changes = []
    for sheet in book:
        for row in sheet.iter_rows(min_row=5):
            for cell in row:
                old = before[sheet.title][cell.row, cell.column]
                if old == cell.value:
                    continue
                if cell.column == flag_columns.get(sheet.title):
                    assert isinstance(old, bool) and cell.value == ('Yes' if old else 'No')
                elif sheet.title == 'signature_comparison' and cell.column in (1, status_column):
                    mapping = methods if cell.column == 1 else statuses
                    assert cell.value == mapping[old]
                else:
                    raise AssertionError((sheet.title, cell.coordinate, old, cell.value))
                changes.append((sheet.title, cell.coordinate))
    assert sum(book['paired_metrics'].cell(r, flag_columns['paired_metrics']).value == 'Yes' for r in range(5, 215)) == 14
    book.save(OUT / 'Supplementary_Data_S3.xlsx')
    s2 = load_workbook(OUT / 'Supplementary_Data_S2.xlsx')
    s2['patient_scores']['A2'] = (
        'Patient-level scores and genomic labels used in Figure 4: SBS3 for HRD, SBS6 for MSI and SBS10 for POLE/POLD1.\n'
        'Exposure is the genomic proportion. Endpoint labels: 0, negative; 1, positive. Blank fields are not applicable or unavailable.\n'
        'Scores are z-standardised within each endpoint only for display in Figure 4c. External and MANTIS sensitivity results are in Data S3.')
    s2['patient_scores'].row_dimensions[2].height = 78
    s2.save(OUT / 'Supplementary_Data_S2.xlsx')
    print(f'Updated notes and labels; {len(changes)} categorical display cells changed; all numerical results preserved.')
    compact_s3()


def compact_s3():
    """Keep the reported results and remove redundant or unused columns only."""
    path = OUT / 'Supplementary_Data_S3.xlsx'
    before = load_workbook(path)
    drop = {
        'paired_metrics': [co + '\n' + field for co in ('TCGA', 'CPTAC')
                           for field in ('Non-zero exposure (n)', 'Spearman P (two-sided)')],
        'tumour_sensitivity': ['Tumour-labelled inputs\nPatients (n)',
                               'Correlation between input-condition scores'],
        'distributions': ['Mean', 'Standard deviation'],
        'mutation_summary': ['Source patients (n)', 'Genomically eligible patients (n)',
                             'Analysed patients (n)'],
        'signature_comparison': ['Nominally replicated\nassociation', 'Paired-target patients (n)',
                                  'Original vs alternative exposures\nPearson r',
                                  'Original vs alternative exposures\nSpearman ρ',
                                  'Original vs alternative exposures\nMean absolute difference'],
    }
    new = Workbook(); new.remove(new.active)
    report = {}
    for sheet in before:
        fields, rows = records(sheet, 4)
        kept = [f for f in fields if f not in drop.get(sheet.title, [])]
        if sheet.title == 'tumour_sensitivity' and 'Tumour-labelled inputs\nPatients (n)' in fields:
            assert all(r['All-slide inputs\nPatients (n)'] == r['Tumour-labelled inputs\nPatients (n)'] for r in rows)
        if sheet.title == 'signature_comparison' and 'Paired-target patients (n)' in fields:
            assert all(r['Matched patients (n)'] == r['Paired-target patients (n)']
                       for r in rows if r['Comparison availability'] != 'Signature unavailable')
        note = sheet['A2'].value
        if sheet.title == 'distributions':
            note = note.replace(' Standard deviations use the population definition.', '')
        elif sheet.title == 'mutation_summary':
            note = ('SNV: single-nucleotide variant. Genomic eligibility requires >50 substitutions.\n'
                    'Mutation counts describe included genomic samples. Patient eligibility and cohort selection are detailed in Data S1.')
        elif sheet.title == 'signature_comparison':
            note = note.replace(' Nominal replication denotes positive associations with P < 0.05 in both cohorts in the primary analysis.', '')
        labels = {f: f for f in kept}
        if sheet.title == 'tumour_sensitivity':
            labels['All-slide inputs\nPatients (n)'] = 'Paired patients (n)'
        result = build(new, sheet.title, sheet['A1'].value, note, kept, rows,
                       labels, sheet.freeze_panes)
        for j, field in enumerate(kept, 1):
            source_col = fields.index(field) + 1
            result.column_dimensions[get_column_letter(j)].width = sheet.column_dimensions[get_column_letter(source_col)].width
            for i in range(4, sheet.max_row + 1):
                src, dst = sheet.cell(i, source_col), result.cell(i, j)
                for attr in ('font', 'fill', 'border', 'alignment', 'protection', 'number_format'):
                    setattr(dst, attr, copy(getattr(src, attr)))
                dst.comment = copy(src.comment)
                assert dst.value == (labels[field] if i == 4 else src.value)
        for i in range(1, sheet.max_row + 1):
            result.row_dimensions[i].height = sheet.row_dimensions[i].height
        report[sheet.title] = {'rows': len(rows), 'columns_before': len(fields),
                               'columns_after': len(kept), 'removed': [f for f in fields if f not in kept]}
    temporary = path.with_suffix('.tmp.xlsx')
    new.save(temporary)
    check = load_workbook(temporary, read_only=True, data_only=True)
    for sheet in before:
        fields = [c.value for c in sheet[4]]
        kept = [f for f in fields if f not in drop.get(sheet.title, [])]
        expected = [[r[fields.index(f)] for f in kept] for r in sheet.iter_rows(min_row=5, values_only=True)]
        assert expected == [list(r) for r in check[sheet.title].iter_rows(min_row=5, values_only=True)]
    check.close()
    temporary.replace(path)
    (ROOT / 'rev/columns_check.json').write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps(report, indent=2))


if __name__=='__main__':
    if '--compact' in sys.argv:
        compact_s3()
    else:
        polish_sheets() if '--polish' in sys.argv else main()
