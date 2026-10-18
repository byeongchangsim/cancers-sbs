"""Reconstruct the original OR endpoint from the complete sequenced profile."""
from pathlib import Path
import json,sys
import numpy as np,pandas as pd
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent.parent;OUT=ROOT/'rev';sys.path.insert(0,str(OUT))
from analyze import binary
source=json.loads((OUT/'cptac_endpoint_source.json').read_text())
rules=pd.read_excel(OUT/'Revision_Data.xlsx',sheet_name='polymerase_rules')
allowed={g:set(f.protein_change.astype(str).str.replace('p.','',regex=False)) for g,f in rules.groupby('gene')}
raw=source['direct_api_full_sequenced_mutation_rows']
labels=pd.read_csv(OUT/'clinical_revision_scores.tsv',sep='\t')
mask=labels.cohort.eq('CPTAC')&labels.endpoint.eq('POLE/POLD1');current=labels.loc[mask].copy()
assert len(current)==82 and set(current.PATIENT)==set(source['current_matched_patient_ids'])
for g in ['POLE','POLD1']:
 events={}
 for row in raw[g]:
  assert row['molecularProfileId']=='ucec_cptac_2020_mutations'
  events.setdefault(row['patientId'],[]).append(row)
 current[g.lower()+'_positive']=current.PATIENT.map(lambda p:int(any(e['proteinChange'].removeprefix('p.') in allowed[g] and e['mutationType']=='Missense_Mutation' for e in events.get(p,[])))) if sys.version_info>=(3,9) else current.PATIENT.map(lambda p:int(any(e['proteinChange'].replace('p.','') in allowed[g] and e['mutationType']=='Missense_Mutation' for e in events.get(p,[]))))
 current[g+'_EVENT_DETAILS']=current.PATIENT.map(lambda p:';'.join(e['proteinChange']+'|'+e['mutationType'] for e in events.get(p,[])) or np.nan)
current['label']=current[['pole_positive','pold1_positive']].max(axis=1).astype(int)
assert int(current.label.sum())==6 and int(current.pold1_positive.sum())==0
changed=labels.loc[mask].loc[labels.loc[mask,'label'].ne(current.label),'PATIENT'].tolist()
assert set(changed).issubset({'C3L-00156'}),changed
labels.loc[mask,current.columns]=current
labels.to_csv(OUT/'clinical_revision_scores.tsv',sep='\t',index=False)
metrics=pd.read_csv(OUT/'clinical_revision.tsv',sep='\t');mi=metrics.cohort.eq('CPTAC')&metrics.endpoint.eq('POLE/POLD1')
result=binary(current.label,current.score)
for k,v in result.items():metrics.loc[mi,k]=v
ix=metrics.index[metrics.cohort.eq('CPTAC')];p=metrics.loc[ix,'p'].to_numpy();order=np.argsort(p);adjusted=np.empty(2);adjusted[order]=np.minimum(1,np.maximum.accumulate(p[order]*[2,1]));metrics.loc[ix,'p_holm_two_endpoints']=adjusted
metrics.to_csv(OUT/'clinical_revision.tsv',sep='\t',index=False)
summary=json.loads((OUT/'summary.json').read_text());summary['clinical']=json.loads(metrics.to_json(orient='records'));summary['polymerase_source_sample_list']='ucec_cptac_2020_sequenced';summary['polymerase_source_correction']='C3L-00156 POLE V411L restored from complete sequenced profile; same original OR rule and image score.';(OUT/'summary.json').write_text(json.dumps(summary,indent=2))
notes=pd.read_csv(OUT/'analysis_notes.tsv',sep='\t')
for i,row in notes.iterrows():
 value=str(row.definition).replace('five positives and 77 negatives','six positives and 76 negatives').replace('5 positive','6 positive').replace('77 negative','76 negative').replace('CPTAC five','CPTAC six').replace('CPTAC5','CPTAC6').replace('all five','five P286R and one V411L')
 notes.loc[i,'definition']=value
key='CPTAC complete mutation-profile audit';value='Original exact-list OR endpoint derived from ucec_cptac_2020_mutations using all 95 sequenced samples; current82 patients all matched. Positive6 (POLE P286R5 plus V411L1), negative76, qualifying POLD1=0. C3L-00156 V411L was absent from the previous all81 sample-list query. The rule and image scores are unchanged; complete source mutation rows and audit are supplied.'
notes=notes[~notes.item.eq(key)];notes=pd.concat([notes,pd.DataFrame([dict(item=key,definition=value)])],ignore_index=True);notes.to_csv(OUT/'analysis_notes.tsv',sep='\t',index=False)
def write_sheet(book,name,frame):
 if name in book:del book[name]
 ws=book.create_sheet(name);ws.append(frame.columns.tolist())
 for row in frame.itertuples(index=False,name=None):ws.append([None if pd.isna(v) else v.item() if isinstance(v,np.generic) else v for v in row])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
book=load_workbook(OUT/'Revision_Data.xlsx')
for name,frame in [('clinical_revision',metrics),('clinical_revision_scores',labels),('analysis_notes',notes),('CPTAC_polymerase_events',pd.read_csv(OUT/'cptac_mutation_evidence.tsv',sep='\t'))]:write_sheet(book,name,frame)
book.save(OUT/'Revision_Data.xlsx')
book=load_workbook(ROOT/'Cancers_evidence.xlsx')
for ws in book:
 for row in ws:
  for cell in row:
   if isinstance(cell.value,str):cell.value=cell.value.replace('five positives and 77 negatives','six positives and 76 negatives').replace('5 positives','6 positives').replace('77 negatives','76 negatives').replace('all five positives carry POLE p.P286R','five POLE p.P286R and one POLE p.V411L')
write_sheet(book,'CPTAC_mutation_evidence',pd.read_csv(OUT/'cptac_mutation_evidence.tsv',sep='\t'))
book.save(ROOT/'Cancers_evidence.xlsx')
(OUT/'corrected_endpoint.json').write_text(json.dumps(dict(changed_patient=changed,rule_unchanged=True,score_unchanged=True,**result,holm_p=float(adjusted[0]),source=source['source_urls']),indent=2))
print(metrics.to_string(index=False))
